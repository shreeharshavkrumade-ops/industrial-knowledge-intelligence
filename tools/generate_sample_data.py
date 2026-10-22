"""Regenerates the synthetic binary documents (PDF, DOCX) and CSV datasets.

All data is FICTIONAL. Random parts use a fixed seed (42) so output is reproducible.
Anchor rows are hand-written so they agree with the narrative sample documents.
Run:  python tools/generate_sample_data.py
"""
from __future__ import annotations

import random
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "data" / "sample_documents"
SEED = 42

ASSETS = {  # asset id -> (name, type)
    "P-101": ("Feedwater Pump", "Pump"), "P-102": ("Cooling Water Pump", "Pump"),
    "M-401": ("Pump Drive Motor", "Motor"), "M-402": ("Compressor Drive Motor", "Motor"),
    "C-202": ("Air Compressor", "Compressor"), "C-205": ("Instrument Air Compressor", "Compressor"),
    "B-301": ("Steam Boiler", "Boiler"), "B-302": ("Standby Steam Boiler", "Boiler"),
}

# fault category -> (component, [description templates], [action templates])
FAULTS = {
    "Bearing failure": ("bearing", [
        "Bearing failure with high vibration and rising bearing temperature",
        "Noisy bearing, grease degraded, bearing worn out",
        "Drive-end bearing damaged, vibration above alarm limit"],
        ["Replaced bearing and re-greased", "Bearing replaced, alignment checked", "Replaced worn bearing and lubricated"]),
    "Overheating": ("cooling fan", [
        "Overheating alarm, discharge temperature high, cooler fouled",
        "High temperature trip, cooling fan slow, oil cooler dusty",
        "Unit overheating during operation, low oil level noted"],
        ["Cleaned cooler and topped up oil", "Cleaned cooler fins and inspected cooling fan", "Replaced cooling fan and cleaned cooler"]),
    "Seal leakage": ("mechanical seal", [
        "Mechanical seal leakage at shaft, water dripping",
        "Seal leak observed, seal faces worn",
        "Leakage from gland area, seal weeping"],
        ["Replaced mechanical seal faces", "Seal replaced and shaft sleeve inspected", "Tightened gland and replaced seal"]),
    "Vibration / misalignment": ("coupling", [
        "Excessive vibration, coupling misalignment suspected",
        "High vibration at coupling end, alignment out of tolerance",
        "Vibration increase after restart, loose coupling bolts"],
        ["Realigned coupling and tightened bolts", "Laser alignment performed, coupling realigned", "Coupling realigned and baseline vibration recorded"]),
    "Cavitation": ("impeller", [
        "Cavitation noise, unstable discharge pressure",
        "Gravel noise from pump, suction pressure low",
        "Cavitation suspected, suction strainer blocked"],
        ["Cleaned suction strainer", "Strainer cleaned and suction level corrected", "Cleaned strainer and inspected impeller"]),
    "Tube fouling / corrosion": ("boiler tube", [
        "Tube fouling with scale deposits, flue gas temperature high",
        "Corrosion found on tubes near header",
        "Scale and fouling reduce heat transfer, high flue gas temperature"],
        ["Chemical flushing of tubes", "Tubes cleaned and corroded area monitored", "Descaled tubes during shutdown"]),
    "Insulation degradation": ("winding", [
        "Low winding insulation resistance, moisture suspected",
        "Insulation resistance below limit at megger test",
        "Winding insulation degraded, terminal box damp"],
        ["Dried winding and retested insulation", "Terminal box cleaned and winding dried", "Winding dried, insulation retested above limit"]),
    "Oil contamination": ("oil filter", [
        "Oil contaminated with particles, oil filter blocked",
        "Oil analysis shows contamination, filter differential pressure high",
        "Dirty compressor oil, filter clogged"],
        ["Oil and oil filter replaced", "Oil changed and filter replaced", "Flushed system and replaced oil filter"]),
}
VALID = {"Pump": ["Bearing failure", "Seal leakage", "Cavitation", "Vibration / misalignment"],
         "Motor": ["Bearing failure", "Overheating", "Insulation degradation", "Vibration / misalignment"],
         "Compressor": ["Overheating", "Oil contamination", "Vibration / misalignment", "Bearing failure"],
         "Boiler": ["Tube fouling / corrosion", "Overheating"]}
DEPT = {"Pump": "Mechanical Maintenance", "Motor": "Electrical Maintenance",
        "Compressor": "Mechanical Maintenance", "Boiler": "Operations"}

# (wo, date, asset, category, description, action, downtime)
ANCHORS = [
    ("WO-2026-0042", "2026-01-20", "P-101", "Routine lubrication", "Routine lubrication of bearings, no abnormalities", "Lubricated bearings with grease EP-2", 0.0),
    ("WO-2026-0101", "2026-03-14", "P-101", "Bearing failure", "Drive-end bearing failure, vibration 7.8 mm/s, bearing temperature 96 C", "Replaced bearing and realigned coupling", 14.0),
    ("WO-2026-0156", "2026-05-02", "P-101", "Seal leakage", "Mechanical seal leakage observed at pump", "Replaced seal faces and inspected shaft sleeve", 6.0),
    ("WO-2026-0203", "2026-06-23", "P-101", "Bearing failure", "Second drive-end bearing failure within four months, misalignment with motor found", "Realigned coupling, replaced and lubricated bearing", 12.0),
    ("WO-2026-0248", "2026-08-19", "P-101", "Cavitation", "Cavitation noise reported, suction strainer blocked", "Cleaned suction strainer", 2.0),
    ("WO-2025-0388", "2025-11-05", "M-401", "Bearing failure", "Non-drive-end bearing overheating and noise", "Replaced bearing and re-greased", 8.0),
    ("WO-2026-0067", "2026-02-10", "M-401", "Bearing failure", "Repeated bearing failure with high vibration", "Replaced bearing", 9.0),
    ("WO-2026-0239", "2026-07-27", "M-401", "Bearing failure", "Third bearing failure in nine months, shaft misalignment found", "Corrected alignment and replaced bearing", 10.0),
    ("WO-2026-0129", "2026-04-17", "C-202", "Overheating", "Overheating alarm, discharge temperature high, oil cooler fouled", "Cleaned oil cooler", 4.0),
    ("WO-2026-0231", "2026-07-08", "C-202", "Overheating", "High discharge temperature trip at 110 C, oil cooler fouled, oil level low", "Cleaned oil cooler, topped up oil, inspected cooling fan", 5.0),
    ("WO-2026-0235", "2026-07-12", "C-202", "Oil contamination", "Oil contaminated with particles after analysis, oil filter blocked", "Oil and oil filter replaced", 3.0),
    ("WO-2026-0251", "2026-08-26", "B-301", "Tube fouling / corrosion", "Tube fouling with scale deposits found at annual inspection", "Chemical flushing of fire-side tubes", 10.0),
]


def build_maintenance(rng: random.Random, nprng: np.random.Generator) -> pd.DataFrame:
    rows = []
    for wo, d, aid, cat, desc, act, dt in ANCHORS:
        rows.append(dict(work_order_id=wo, date=d, equipment_id=aid, equipment_type=ASSETS[aid][1], component=FAULTS.get(cat, ("bearing",))[0],
                         fault_description=desc, fault_label=cat, action_taken=act, downtime_hours=dt, department=DEPT[ASSETS[aid][1]]))
    # weighted generation so P-101 / M-401 bearings, C-202 overheating recur (patterns to discover)
    weights = {"P-101": 4, "P-102": 2, "M-401": 3, "M-402": 2, "C-202": 4, "C-205": 2, "B-301": 3, "B-302": 2}
    ids = list(weights)
    start = date(2025, 10, 1)
    n = 0
    while n < 108:
        aid = rng.choices(ids, weights=[weights[i] for i in ids])[0]
        etype = ASSETS[aid][1]
        cats = VALID[etype]
        if aid == "P-101":
            cats = ["Bearing failure"] * 3 + cats
        if aid == "C-202":
            cats = ["Overheating"] * 3 + cats
        if aid == "M-401":
            cats = ["Bearing failure"] * 3 + cats
        cat = rng.choice(cats)
        comp, descs, acts = FAULTS[cat]
        d = start + timedelta(days=rng.randint(0, 330))
        rows.append(dict(work_order_id=f"WO-{d.year}-{900 + n:04d}", date=d.isoformat(), equipment_id=aid, equipment_type=etype,
                         component=comp, fault_description=rng.choice(descs), fault_label=cat, action_taken=rng.choice(acts),
                         downtime_hours=round(float(nprng.gamma(2.0, 2.5)), 1), department=DEPT[etype]))
        n += 1
    df = pd.DataFrame(rows).sort_values("date").reset_index(drop=True)
    # Pre-event condition features + synthetic target. The target is generated from a noisy logistic rule,
    # so any classifier can only recover that synthetic rule - NOT real failure physics.
    df["operating_hours_since_overhaul"] = nprng.integers(200, 9000, len(df))
    df["days_since_last_pm"] = nprng.integers(5, 200, len(df))
    df["vibration_mm_s"] = np.round(nprng.gamma(4.0, 1.0, len(df)), 2)
    df["temperature_c"] = np.round(nprng.normal(72, 12, len(df)), 1)
    z = (-3.9 + 0.00018 * df.operating_hours_since_overhaul + 0.012 * df.days_since_last_pm
         + 0.35 * df.vibration_mm_s + 0.04 * (df.temperature_c - 70) + nprng.normal(0, 0.9, len(df)))
    df["failure_within_30d"] = (1 / (1 + np.exp(-z)) > 0.5).astype(int)
    df["severity"] = pd.cut(df.downtime_hours, [-1, 3, 8, 1e9], labels=["Low", "Medium", "High"]).astype(str)
    # a few deliberately messy values to exercise cleaning code
    df.loc[[7, 40], "downtime_hours"] = np.nan
    df.loc[[15], "component"] = ""
    return df


INSPECTIONS = [
    ("INS-001", "2026-05-18", "P-102", "Vibration Analysis", "Inspection", "Vibration 3.1 mm/s within limits; minor flange gasket leak tightened", "Pass"),
    ("INS-002", "2026-07-10", "P-101", "Vibration Analysis", "Mechanical Maintenance", "Vibration 5.2 mm/s above 4.5 mm/s limit at drive end", "Observation"),
    ("INS-003", "2026-09-05", "M-401", "Vibration Analysis", "Electrical Maintenance", "Vibration 7.1 mm/s, bearing noise audible", "Fail"),
    ("INS-004", "2026-03-20", "M-401", "Insulation Resistance Test", "Electrical Maintenance", "Insulation resistance 48 MOhm acceptable", "Pass"),
    ("INS-005", "2026-06-15", "M-402", "Insulation Resistance Test", "Electrical Maintenance", "Insulation resistance 55 MOhm acceptable", "Pass"),
    ("INS-006", "2026-08-12", "B-301", "Boiler Safety Inspection", "Inspection", "Tube fouling, minor economiser corrosion, thermocouple drift of 6 C", "Observation"),
    ("INS-007", "2025-12-01", "B-302", "Boiler Safety Inspection", "Inspection", "No defects found on standby boiler", "Pass"),
    ("INS-008", "2026-01-15", "B-301", "Relief Valve Test", "Inspection", "Relief valve lifted within tolerance", "Pass"),
    ("INS-009", "2025-08-20", "C-202", "Relief Valve Test", "Mechanical Maintenance", "Relief valve lifted within tolerance", "Pass"),
    ("INS-010", "2026-07-09", "C-202", "Oil Analysis", "Mechanical Maintenance", "Oil contaminated with particles above limit", "Fail"),
    ("INS-011", "2026-02-01", "C-205", "Oil Analysis", "Mechanical Maintenance", "Oil condition acceptable", "Pass"),
    ("INS-012", "2026-06-02", "P-101", "Oil Analysis", "Mechanical Maintenance", "Bearing grease sample acceptable", "Pass"),
]
FREQ = {"Vibration Analysis": 90, "Insulation Resistance Test": 180, "Boiler Safety Inspection": 180, "Relief Valve Test": 365, "Oil Analysis": 120}


def build_inspections() -> pd.DataFrame:
    rows = []
    for iid, d, aid, typ, dept, finding, status in INSPECTIONS:
        due = (date.fromisoformat(d) + timedelta(days=FREQ[typ])).isoformat()
        rows.append(dict(inspection_id=iid, date=d, equipment_id=aid, inspection_type=typ, department=dept,
                         finding=finding, status=status, next_due_date=due))
    return pd.DataFrame(rows)


# FICTIONAL requirements, for illustration only; not legal advice.
REQUIREMENTS = [
    ("REQ-01", "Pump vibration analysis", "Vibration analysis of each pump shall be recorded at least every 90 days.", "periodic_inspection", "Pump", "Vibration Analysis", 90, ""),
    ("REQ-02", "Motor vibration analysis", "Vibration analysis of each motor shall be recorded at least every 90 days.", "periodic_inspection", "Motor", "Vibration Analysis", 90, ""),
    ("REQ-03", "Motor insulation test", "Winding insulation resistance of each motor shall be tested every 180 days.", "periodic_inspection", "Motor", "Insulation Resistance Test", 180, ""),
    ("REQ-04", "Boiler safety inspection", "Each boiler shall have a documented safety inspection every 180 days.", "periodic_inspection", "Boiler", "Boiler Safety Inspection", 180, ""),
    ("REQ-05", "Relief valve test", "Pressure relief valves on boilers and compressors shall be tested every 365 days.", "periodic_inspection", "Boiler;Compressor", "Relief Valve Test", 365, ""),
    ("REQ-06", "Compressor oil analysis", "Compressor oil shall be analysed every 120 days.", "periodic_inspection", "Compressor", "Oil Analysis", 120, ""),
    ("REQ-07", "Failed inspection follow-up", "A failed inspection shall be followed by a work order within 14 days.", "follow_up", "All", "", 14, ""),
    ("REQ-08", "Lockout procedure documented", "A documented lockout/tagout procedure shall exist before motor inspection.", "documented_procedure", "All", "", 0, "lock out;lockout;tagout;isolate"),
    ("REQ-09", "Boiler water treatment documented", "A boiler water treatment procedure shall be documented.", "documented_procedure", "All", "", 0, "water treatment;blowdown;TDS"),
]


def make_pdf() -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(DOCS / "boiler_b301_inspection_report.pdf"), pagesize=A4)
    def page(lines):
        y = 800
        for i, ln in enumerate(lines):
            c.setFont("Helvetica-Bold" if i == 0 else "Helvetica", 14 if i == 0 else 10.5)
            c.drawString(50, y, ln); y -= 24 if i == 0 else 17
        c.showPage()
    page(["BOILER B-301 ANNUAL INSPECTION REPORT", "Inspection date: 2026-08-12", "Inspector: Inspection Department",
          "Scope: shell, tubes, burner, pressure relief valve and flue gas thermocouple.", "",
          "Findings:",
          "1. Tube fouling and scale deposits were found on the fire-side tubes of Boiler B-301,",
          "    reducing heat transfer efficiency.",
          "2. Minor corrosion was observed on the economiser tubes near the lower header.",
          "3. The pressure relief valve lifted at 10.2 bar against a set pressure of 10.0 bar,",
          "    which is within tolerance.",
          "4. The flue gas thermocouple showed reading drift of 6 C."])
    page(["RECOMMENDATIONS - BOILER B-301", "",
          "- Clean the fouled tubes by chemical flushing during the next shutdown.",
          "- Calibrate or replace the flue gas thermocouple.",
          "- Re-inspect the corroded economiser area within 90 days.", "",
          "Compliance: this inspection was performed in accordance with IKI-REG-12.",
          "Next boiler safety inspection due: 2027-02-08.",
          "Prepared by the Inspection Department."])
    c.save()


def make_docx() -> None:
    from docx import Document
    d = Document()
    d.add_heading("SAFETY PROCEDURE - ELECTRIC MOTOR INSPECTION", 1)
    d.add_paragraph("Document No: IKI-SOP-MOT-01")
    d.add_paragraph("Department: Electrical Maintenance")
    d.add_paragraph("Purpose: This procedure applies to inspection of Motor M-401 and Motor M-402.")
    d.add_heading("Safety precautions before motor inspection", 2)
    for s in ["Notify the control room and obtain a work permit.",
              "Switch off and lock out the motor at the motor control centre; apply a personal lock and tag (lockout/tagout) in line with IKI-SAF-007.",
              "Verify zero voltage with a tested voltage detector.",
              "Wait for the rotor to stop completely and discharge capacitors.",
              "Wear insulated gloves, safety glasses and arc-rated clothing.",
              "Never remove the coupling guard while the motor is rotating."]:
        d.add_paragraph(s, style="List Number")
    d.add_heading("Inspection steps", 2)
    d.add_paragraph("Check winding insulation resistance with a megger. Insulation resistance below 5 MOhm requires winding drying or rewinding assessment. "
                    "Check bearing noise and bearing temperature. Inspect the terminal box for loose connections and overheating marks.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Hazard", "Control"
    for h, c in [("Stored electrical energy", "Lockout and verify zero voltage"), ("Rotating parts", "Keep coupling guard fitted"), ("Hot surfaces", "Allow motor to cool, wear gloves")]:
        r = t.add_row().cells; r[0].text, r[1].text = h, c
    d.save(str(DOCS / "motor_inspection_safety_procedure.docx"))


def main() -> None:
    rng, nprng = random.Random(SEED), np.random.default_rng(SEED)
    make_pdf(); make_docx()
    build_maintenance(rng, nprng).to_csv(ROOT / "data" / "maintenance_records.csv", index=False)
    build_inspections().to_csv(ROOT / "data" / "inspection_records.csv", index=False)
    pd.DataFrame(REQUIREMENTS, columns=["requirement_id", "title", "requirement_text", "check_type", "applies_to",
                                         "inspection_type", "max_interval_days", "evidence_keywords"]).to_csv(ROOT / "data" / "compliance_requirements.csv", index=False)
    print("sample data written")


if __name__ == "__main__":
    main()
